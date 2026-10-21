"""Read Devon's price list straight out of the official Devon calculator (.xlsm).

The calculator's "Master" sheet lists every Devon product with:
  A = product number, B = description, C = quantity (blank in the template),
  D = list price (usually =INDEX(L:N, $J$10), i.e. picks a fabric-grade column),
  I = weight, L/M/N = list price for Keystone / Summit / Aspen.
Net price is list x the net multiplier in F18.

Reading the catalog from the calculator itself means a new Devon calculator can be
dropped into templates/ without touching any code.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

GRADES = ("Keystone", "Summit", "Aspen")

_LREF = re.compile(r"^=\s*L(\d+)\s*(?:\+\s*([\d.]+))?\s*$")
_RATE = re.compile(r"\*\s*([\d.]+)\s*$")


@dataclass
class DevonItem:
    part: str
    description: str
    row: int
    section: str
    weight: float
    grade_prices: tuple[float, float, float] | None  # None when D holds a fixed price
    fixed_price: float | None

    def list_price(self, grade: int) -> float:
        """List price for grade index 0=Keystone, 1=Summit, 2=Aspen."""
        if self.grade_prices is None:
            return self.fixed_price or 0.0
        return self.grade_prices[grade]


@dataclass
class Catalog:
    items: dict[str, DevonItem]
    net_multiplier: float
    surcharge_rate: float
    title: str
    first_row: int
    last_row: int
    sections: list[str] = field(default_factory=list)

    def __contains__(self, part: str) -> bool:
        return part in self.items

    def get(self, part: str) -> DevonItem | None:
        return self.items.get(part)


def _num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _grade_col(value, l_price: float | None) -> float | None:
    """Evaluate the M/N cells, which are numbers or simple '=L26+36' formulas."""
    n = _num(value)
    if n is not None:
        return n
    if isinstance(value, str) and l_price is not None:
        m = _LREF.match(value)
        if m:
            return l_price + float(m.group(2) or 0)
    return l_price


def _is_formula(v) -> bool:
    return hasattr(v, "text") or (isinstance(v, str) and v.startswith("="))


def load_catalog(path: str | Path) -> Catalog:
    wb = openpyxl.load_workbook(path)
    ws = wb["Master"]

    net_multiplier = _num(ws["F18"].value) or 0.3
    title = str(ws["A3"].value or "")

    # The product list runs from the header row ("PRODUCT NUMBER") to "SUB-TOTAL: PRODUCT ONLY".
    first_row = last_row = None
    surcharge_rate = 0.0
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a == "PRODUCT NUMBER" and r > 20:
            first_row = r + 1
        elif isinstance(a, str) and a.startswith("SUB-TOTAL: PRODUCT ONLY"):
            last_row = r - 1
        elif a == "SURCHARGE":
            m = _RATE.search(str(ws.cell(r, 7).value or ""))
            surcharge_rate = float(m.group(1)) if m else 0.0
    if first_row is None or last_row is None:
        raise ValueError(f"{path}: could not find the product list on the Master sheet")

    items: dict[str, DevonItem] = {}
    sections: list[str] = []
    section = ""
    for r in range(first_row, last_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not isinstance(a, str) or not a.strip():
            continue
        a = a.strip()
        if a.startswith(("-->>", ">>")):
            continue
        if not (isinstance(b, str) and b.strip()):
            section = a
            sections.append(a)
            continue
        if a in items:  # e.g. ACC-110 is listed twice; the first row is the no-charge one
            continue
        d = ws.cell(r, 4).value
        l_price = _num(ws.cell(r, 12).value)
        if _is_formula(d) and l_price is not None:
            prices = (
                l_price,
                _grade_col(ws.cell(r, 13).value, l_price),
                _grade_col(ws.cell(r, 14).value, l_price),
            )
            fixed = None
        else:
            prices, fixed = None, _num(d) or 0.0
        items[a] = DevonItem(
            part=a,
            description=b.strip(),
            row=r,
            section=section,
            weight=_num(ws.cell(r, 9).value) or 0.0,
            grade_prices=prices,
            fixed_price=fixed,
        )

    return Catalog(
        items=items,
        net_multiplier=net_multiplier,
        surcharge_rate=surcharge_rate,
        title=title,
        first_row=first_row,
        last_row=last_row,
        sections=sections,
    )
