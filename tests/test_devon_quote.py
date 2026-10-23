import zipfile
from pathlib import Path

import openpyxl
import pytest

from devon_quote.bom import BomLine, load_bom, parse_table
from devon_quote.catalog import load_catalog
from devon_quote.mapper import Mapper, Override, devon_height
from devon_quote.pricing import estimate
from devon_quote.writer import QuoteHeader, write_quote

ROOT = Path(__file__).resolve().parent.parent
from devon_quote.service import find_template  # noqa: E402

TEMPLATE = find_template()
SAMPLE = ROOT / "tests" / "fixtures" / "sample_ais_divi.sif"


@pytest.fixture(scope="module")
def catalog():
    if TEMPLATE is None:
        pytest.skip("put the Devon calculator .xlsm at the top of the repo")
    return load_catalog(TEMPLATE)


def line(part, desc="", qty=1):
    return BomLine(part=part, description=desc, qty=qty)


def test_catalog_prices_by_grade(catalog):
    item = catalog.get("HM01.5348")
    assert item.list_price(0) == pytest.approx(425.04)  # Keystone
    assert item.list_price(1) == pytest.approx(491.04)  # Summit = L + 66
    assert item.list_price(2) == pytest.approx(591.04)  # Aspen = L + 166
    assert catalog.net_multiplier == pytest.approx(0.3)
    assert catalog.surcharge_rate == pytest.approx(0.09)
    # 85" glass-stacker panels carry a fixed price regardless of grade
    assert catalog.get("HM51.8548").list_price(2) == catalog.get("HM51.8548").list_price(0)


def test_parse_sif():
    lines = load_bom(SAMPLE)
    assert len(lines) == 24
    first = lines[0]
    assert first.part == "E-UN8W4BIFS" and first.qty == 5
    panel = next(l for l in lines if l.part == "P-DLFRPM5024")
    assert ("Fabric Side A", "New England Night") in [(o.group, o.description) for o in panel.options]


def test_parse_table():
    rows = [["Project X"], ["Part Number", "Description", "Qty"], ["W-WS2460", "Worksurface", "4"], ["", "", ""]]
    assert [(l.part, l.qty) for l in parse_table(rows)] == [("W-WS2460", 4)]


def test_heights():
    assert devon_height(50) == (53, "")
    assert devon_height(42) == (39, "")
    assert devon_height(82) == (85, "")
    assert devon_height(66) == (67, "")
    assert devon_height(30)[0] == 39 and devon_height(30)[1]
    assert devon_height(90)[0] == 85


@pytest.mark.parametrize(
    "part,desc,expected",
    [
        ("P-DLFRPM5024", "Divi Linear Panel - Hard Surf Perm/Mono - NPWR - 50H x 24W", "HM01.5324"),
        ("P-DLFPPM5060", "Divi Linear Panel - Hard Surf Perm/Mono - PWR - 50H x 60W", "HM01.5360.P4"),
        ("P-DLQGWRMM5042", "Divi Linear Qtr Glass Panel - AMono/BMono - NPWR - 50H x 42W", "HM51.5342"),
        ("P-DLQGWPMM5060", "Divi Linear Qrt Glass Panel - AMono/BMono - PWR - 50H x 60W", "HM51.5360.P4"),
        ("P-DLPF350", "Divi Linear Post Filler - 3 Way - 50H", "HMA137.53"),
        ("P-DLPF450", "Divi Linear Post Filler - 4 Way - 50H", "HMC132.48"),
        ("W-WS2442", "Rectangular - 24D x 42W", "WS2442"),
        ("WPS-UNCT24S", "Shared Cantilever - Single - 24D", "HMP226.22"),
        ("WPS-VC48", "U-Channel - 48W Black", "WPS-VC48"),
        ("E-MW8W4DO2", "Duplex Outlet - Line 2 - Black", "BE-REC2"),
    ],
)
def test_rules(catalog, part, desc, expected):
    m = Mapper(catalog).map_line(line(part, desc))
    assert m.targets[0][0] == expected
    assert m.status in ("ok", "review")


def test_two_way_post_is_a_corner(catalog):
    # AIS lists no part for straight joints, so a 2-way post is always a corner
    m = Mapper(catalog).map_line(line("P-DLPF250", "Divi Linear Post Filler - 2 -Way 50H"))
    assert m.targets == [("HMA130.53", 1)] and m.status == "ok"


def test_no_2d_warns_that_panel_to_panel_joints_are_missing(catalog):
    res = Mapper(catalog).map([line("P-DLFRPM5024", "Divi Linear Panel - NPWR - 50H x 24W", 2)])
    assert any("doesn't list straight panel-to-panel joints" in w for w in res.warnings)


def test_cantilevers_become_pairs(catalog):
    res = Mapper(catalog).map([
        line("WPS-UNCT24L", "Regular Cantilever - Left Hand - 24D", 3),
        line("WPS-UNCT24R", "Regular Cantilever - Right Hand - 24D", 3),
    ])
    assert res.quantities["HMP225.22"] == 3
    assert res.quantities["ACC-107"] == 3
    assert not res.warnings


@pytest.mark.parametrize("lefts,rights,pairs", [(15, 15, 15), (7, 8, 8), (3, 0, 2), (1, 0, 1), (0, 5, 3)])
def test_cantilever_pairs_are_total_over_two_rounded_up(catalog, lefts, rights, pairs):
    lines = []
    if lefts:
        lines.append(line("WPS-UNCT24L", "Regular Cantilever - Left Hand - 24D", lefts))
    if rights:
        lines.append(line("WPS-UNCT24R", "Regular Cantilever - Right Hand - 24D", rights))
    res = Mapper(catalog).map(lines)
    assert res.quantities["HMP225.22"] == pairs
    assert not res.warnings  # unequal lefts and rights are fine: the rule covers them


def test_cantilevers_recognized_by_description(catalog):
    res = Mapper(catalog).map([
        line("X-CANT-L", "Regular Cantilever - Left Hand - 16D", 2),
        line("X-CANT-R", "Regular Cantilever - Right Hand - 16D", 3),
        line("X-CANT-S", "Shared Cantilever - Single - 24D", 4),
    ])
    assert res.quantities["HMP165.16"] == 3 and res.quantities["HMP226.22"] == 4
    assert not res.unmapped


def test_odd_worksurface_gets_custom_cut(catalog):
    m = Mapper(catalog).map_line(line("W-WS2454", "Rectangular - 24D x 54W"))
    assert m.targets == [("WS2460", 1), ("MIS-002", 1)] and m.status == "review"


def test_unknown_part_is_reported(catalog):
    res = Mapper(catalog).map([line("X-MYSTERY", "Something", 2)])
    assert res.unmapped and not res.quantities


def test_override_wins(catalog):
    mapper = Mapper(catalog, [Override("P-DLPF250", "HMP213.48", 1, "straight runs only")])
    m = mapper.map_line(line("P-DLPF250", "Divi Linear Post Filler - 2 -Way 50H"))
    assert m.targets == [("HMP213.48", 1)] and m.status == "ok"
    skip = Mapper(catalog, [Override("E-MX*", "", 1, "")]).map_line(line("E-MX8W4PTPC21", "Jumper"))
    assert skip.targets == [] and skip.status == "ok"


def test_sample_maps_cleanly(catalog):
    res = Mapper(catalog).map(load_bom(SAMPLE))
    assert not res.unmapped
    assert res.quantities["HMA128.53"] == 23  # end-of-run shrouds -> finished ends at 53"H
    assert all(p in catalog for p in res.quantities)


def test_writer_fills_calculator(tmp_path, catalog):
    res = Mapper(catalog).map(load_bom(SAMPLE))
    out = write_quote(TEMPLATE, tmp_path / "q.xlsm", catalog, res.quantities,
                      QuoteHeader(project="Test & Co", ship_to="33010", grade=1, delivery=1000, stations=10))
    assert "xl/vbaProject.bin" in zipfile.ZipFile(out).namelist()
    ws = openpyxl.load_workbook(out, keep_vba=True)["Master"]
    filled = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(catalog.first_row, catalog.last_row + 1) if ws.cell(r, 3).value}
    assert filled == res.quantities
    assert ws["B6"].value == "Test & Co" and ws["J10"].value == 2 and ws["G444"].value == 1000
    assert ws["B4"].value is None  # Devon's quote number is cleared
    assert str(ws["G442"].value).startswith("=")  # formulas untouched

    est = estimate(catalog, res.quantities, 1, 1000, 10)
    assert est.total == pytest.approx(est.product_net * 1.09 + 1000)


def test_delivery_estimate_from_weight():
    from devon_quote.pricing import estimate_delivery

    # $445 + $0.27/lb, rounded to $5: fitted to Devon's freight on 25 quotes to OF4S
    assert estimate_delivery(0) == 0
    assert estimate_delivery(313.5) == 530  # Devon quoted $426.85
    assert estimate_delivery(2238.24) == 1050  # Devon quoted $823.17
    assert estimate_delivery(14823.645) == 4445  # Devon quoted $4,500


def test_blank_delivery_is_estimated(catalog):
    from devon_quote.pricing import estimate, estimate_delivery

    q = {"HMA130.53": 4}
    auto = estimate(catalog, q, 0, None, None)
    assert auto.delivery_auto and auto.delivery == estimate_delivery(auto.weight)
    assert auto.total == pytest.approx(auto.product_net + auto.surcharge + auto.delivery)
    typed = estimate(catalog, q, 0, 0, None)
    assert not typed.delivery_auto and typed.delivery == 0  # an explicit $0 (pickup) stays $0


@pytest.mark.parametrize("part,devon", [("P-DLPFPP250", "HMA130.53"), ("P-DLPFPP366", "HMA137.67"),
                                        ("P-DLPF250", "HMA130.53"), ("P-DLPF366", "HMA137.67")])
def test_pp_post_fillers_map_like_plain_ones(catalog, part, devon):
    m = Mapper(catalog).map_line(line(part, "Divi Linear Post Filler"))
    assert m.targets == [(devon, 1)] and m.status == "ok"


def test_divi50_choice(catalog):
    panel = line("P-DLFRPM5024", "Divi Linear Panel - NPWR - 50H x 24W")
    post = line("P-DLPF250", "Divi Linear Post Filler - 2 -Way 50H")
    assert Mapper(catalog).map_line(panel).targets == [("HM01.5324", 1)]  # default 53"
    m47 = Mapper(catalog, divi50=47)
    assert m47.map_line(panel).targets == [("HM01.4724", 1)]
    assert m47.map_line(post).targets == [("HMA130.47", 1)]  # connectors follow the panel height
    with pytest.raises(ValueError):
        Mapper(catalog, divi50=50)
