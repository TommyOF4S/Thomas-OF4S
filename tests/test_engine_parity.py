import zipfile
"""The hosted page runs a JavaScript port of the engine. This checks that it produces
exactly the same quote as the Python package (needs Node and the jszip npm package;
set NODE_PATH to where jszip is installed)."""

import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from devon_quote.bom import load_bom
from devon_quote.catalog import load_catalog
from devon_quote.mapper import Mapper, Override
from devon_quote.pricing import estimate
from devon_quote.report import write_crosswalk, write_report
from devon_quote.service import find_template

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "tests" / "fixtures" / "sample_ais_divi.sif"
HARNESS = ROOT / "tests" / "engine_harness.js"
TEMPLATE = find_template()


def _node_ok() -> bool:
    if not shutil.which("node"):
        return False
    return subprocess.run(["node", "-e", "require('jszip')"], capture_output=True).returncode == 0


pytestmark = [
    pytest.mark.skipif(TEMPLATE is None, reason="put the Devon calculator .xlsm in templates/"),
    pytest.mark.skipif(not _node_ok(), reason="needs node with jszip on NODE_PATH"),
]


def run_js(tmp_path, bom, out_name="q.xlsx", **opts):
    out = tmp_path / out_name
    p = subprocess.run(["node", str(HARNESS), str(TEMPLATE), str(bom), str(out), json.dumps(opts)],
                       capture_output=True, text=True, check=True)
    return json.loads(p.stdout), out


@pytest.mark.parametrize("grade", [0, 1, 2])
def test_same_quote_as_python(tmp_path, grade):
    js, _ = run_js(tmp_path, SAMPLE, grade=grade, delivery=4500, stations=30, project="Parity")
    cat = load_catalog(TEMPLATE)

    # Catalog: every part, row, grade price and weight
    py_cat = {p: [i.row, *(i.list_price(g) for g in range(3)), i.weight] for p, i in cat.items.items()}
    assert list(js["catalog"]) == list(py_cat)
    for p, vals in py_cat.items():
        assert js["catalog"][p] == pytest.approx(vals), p
    assert js["netMultiplier"] == cat.net_multiplier and js["surchargeRate"] == cat.surcharge_rate

    lines = load_bom(SAMPLE)
    res = Mapper(cat).map(lines)
    assert js["quantities"] == res.quantities
    assert [m[:3] for m in js["mappings"]] == [[m.source.part, m.status, m.note] for m in res.mappings]
    assert [[[t, r] for t, r in m.targets] for m in res.mappings] == [m[3] for m in js["mappings"]]
    est = estimate(cat, res.quantities, grade, 4500, 30)
    assert js["total"] == pytest.approx(est.total) and js["product"] == pytest.approx(est.product_net)

    cw = tmp_path / "cw.csv"
    write_crosswalk(cw, res)
    assert js["crosswalk"] == cw.read_bytes().decode("utf-8")

    rp = tmp_path / "r.md"
    write_report(rp, project="Parity", bom_file=SAMPLE.name, quote_file="q.xlsx", catalog=cat, grade=grade,
                 lines=lines, result=res, est=est)
    strip = lambda t: [l for l in t.splitlines() if "overrides.csv" not in l and "save a correction" not in l]
    assert strip(js["report"]) == strip(rp.read_text(encoding="utf-8"))


def test_overrides_and_edge_parts(tmp_path):
    bom = tmp_path / "edge.csv"
    bom.write_text(
        "Part Number,Description,Qty\n"
        "P-DLPF250,Divi Linear Post Filler - 2 -Way 50H,4\n"
        "W-WS2454,Rectangular - 24D x 54W,2\n"
        "P-DLFRPM3024,Divi Linear Panel - NPWR - 30H x 24W,3\n"
        "WPS-UNCT24L,Regular Cantilever - Left Hand - 24D,3\n"
        "WPS-UNCT24R,Regular Cantilever - Right Hand - 24D,4\n"
        "X-CANT-L,Regular Cantilever - Left Hand - 16D,1\n"
        "E-MX8W4PTPC21,21 in. Electrical Jumper,2\n"
        "X-MYSTERY,\"Something, odd\",1\n"
    )
    ovs = [{"ais_part": "P-DLPF250", "devon_part": "HMP213.48", "ratio": 1, "note": "straight runs"},
           {"ais_part": "E-MX*", "devon_part": "", "ratio": 1, "note": ""}]
    js, _ = run_js(tmp_path, bom, overrides=ovs)
    cat = load_catalog(TEMPLATE)
    res = Mapper(cat, [Override(o["ais_part"].upper(), o["devon_part"], o["ratio"], o["note"]) for o in ovs]).map(load_bom(bom))
    assert js["quantities"] == res.quantities
    assert [m[:3] for m in js["mappings"]] == [[m.source.part, m.status, m.note] for m in res.mappings]
    assert js["warnings"] == res.warnings


def test_xlsx_output_is_clean(tmp_path):
    js, out = run_js(tmp_path, SAMPLE, grade=1, delivery=1000, stations=10, project="Hosted")
    wb = openpyxl.load_workbook(out)  # loads as a plain workbook, no macros
    ws = wb["Master"]
    filled = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(26, 437) if ws.cell(r, 3).value}
    assert filled == js["quantities"]
    assert ws["B6"].value == "Hosted" and ws["J10"].value == 2 and ws["G444"].value == 1000
    import zipfile
    z = zipfile.ZipFile(out)
    assert "xl/vbaProject.bin" not in z.namelist()
    assert b"macroEnabled" not in z.read("[Content_Types].xml")


def test_delivery_estimated_from_weight_matches_python(tmp_path):
    js, out = run_js(tmp_path, SAMPLE, stations=30, project="Auto")  # no delivery: estimated
    cat = load_catalog(TEMPLATE)
    lines = load_bom(SAMPLE)
    res = Mapper(cat).map(lines)
    est = estimate(cat, res.quantities, 0, None, 30)
    assert est.delivery_auto and est.delivery > 0
    assert js["total"] == pytest.approx(est.total)
    rp = tmp_path / "r.md"
    write_report(rp, project="Auto", bom_file=SAMPLE.name, quote_file="q.xlsx", catalog=cat, grade=0,
                 lines=lines, result=res, est=est)
    strip = lambda t: [l for l in t.splitlines() if "overrides.csv" not in l and "save a correction" not in l]
    assert strip(js["report"]) == strip(rp.read_text(encoding="utf-8"))
    assert any(l.startswith("| Delivery (estimated from ") for l in strip(js["report"]))
    # the estimate goes into the calculator's delivery cell
    sheet = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode()
    assert f'r="G444"' in sheet and f">{est.delivery:g}<" in sheet.split('r="G444"', 1)[1][:80]


def test_pp_post_fillers_match_python(tmp_path):
    bom = tmp_path / "pp.csv"
    bom.write_text("Part Number,Description,Qty\nP-DLFRPM5024,Divi Linear Panel - NPWR - 50H x 24W,4\n"
                   "P-DLPFPP250,Divi Linear Post Filler PP - 2 Way - 50H,3\nP-DLPFPP350,Divi Linear Post Filler PP - 3 Way - 50H,2\n"
                   "P-DLPF250,Divi Linear Post Filler - 2 Way - 50H,1\n", encoding="utf-8")
    js, _ = run_js(tmp_path, bom, delivery=0)
    res = Mapper(load_catalog(TEMPLATE)).map(load_bom(bom))
    assert js["quantities"] == res.quantities
    assert res.quantities["HMA130.53"] == 4 and res.quantities["HMA137.53"] == 2
